import os
import shutil
import zipfile
import docx

def update_docx():
    docx_path = r"docs/Arena3-SRS.docx"
    backup_path = r"docs/Arena3-SRS.docx.bak"
    if not os.path.exists(backup_path):
        shutil.copyfile(docx_path, backup_path)
        print("Created backup at:", backup_path)

    doc = docx.Document(docx_path)

    # 1. Update Title / Version on Cover Page
    for p in doc.paragraphs[:20]:
        if "Version 1.6" in p.text:
            p.text = p.text.replace("Version 1.6", "Version 1.7")
            print("Updated Cover Version to 1.7")
        if "5 October 2026" in p.text:
            p.text = p.text.replace("5 October 2026", "11 October 2026")
            print("Updated Cover Date to 11 October 2026")

    # Update Headers and P6
    for s in doc.sections:
        for p in s.header.paragraphs:
            if "Version 1.6" in p.text:
                p.text = p.text.replace("Version 1.6", "Version 1.7")
                print("Updated Section Header to Version 1.7")
    if len(doc.paragraphs) > 6 and "version 1.6" in doc.paragraphs[6].text:
        doc.paragraphs[6].text = doc.paragraphs[6].text.replace("version 1.6", "version 1.7")
        print("Updated Paragraph 6 to version 1.7")

    # 2. Update Table 0 (Revision History)
    t0 = doc.tables[0]
    has_1_7 = any("1.7" in cell.text for row in t0.rows for cell in row.cells)
    if not has_1_7:
        new_row = t0.add_row()
        new_row.cells[0].text = "1.7"
        new_row.cells[1].text = "11 October 2026"
        new_row.cells[2].text = (
            "Updated as-built to FPT University Capstone / SWR302 English standard: 9-actor model "
            "(Guest, Member, Receptionist, Coach, Manager, AI Assistant, Mail Service, Scheduler, payOS), "
            "DFD Level 0 Context Diagram with 100% English noun phrase data flows, UML 2.0 Swimlane Register "
            "with Mail OTP across 5 partitions, VietQR dynamic payment links, and verified attendance workflows."
        )
        print("Added Version 1.7 to Revision History table")

    # 3. Update Paragraph 48 (Context Diagram description)
    p48 = doc.paragraphs[48]
    if "Figure 1 puts Arena3" in p48.text:
        p48.text = (
            "Figure 1 puts Arena3 in its operational setting (DFD Level 0 Context Diagram). "
            "The system interacts with seven external entities and actors: Guest, Member, Receptionist, "
            "Coach, Manager, payOS Payment Gateway, and Mail Delivery Service. All data flows are strictly "
            "noun-based phrases compliant with FPT University SE curriculum standards (e.g., Account Registration Data, "
            "Schedule Inquiries, Trainee Performance Evaluations, Booking Invoices, and Webhook Payloads). "
            "The background scheduler executes automated routines (expiring holds, marking no-shows, "
            "generating classes). payOS provides live VietQR dynamic payment links and confirms payments via signed webhooks."
        )
        print("Updated Paragraph 48 (Context Diagram description)")

    # 4. Update Paragraph 52 (Use Case Diagram description)
    p52 = doc.paragraphs[52]
    if "Figure 2 draws the system boundary" in p52.text:
        p52.text = (
            "Figure 2 illustrates the system boundary, the 9 actors outside it, and the functional use cases "
            "organized across six core business packages compliant with UML 2.0. Solid lines connect primary "
            "initiating actors. Dashed lines indicate <<include>> and <<extend>> dependencies. Guest browses public "
            "schedules and registers an account. Member books courts, joins classes, purchases plans, views training "
            "progress, and queries the AI Assistant. Receptionist manages till shifts, processes cash and VietQR POS "
            "sales, checks members in via QR pass or member code, and handles gear rentals. Coach manages assigned "
            "classes, attendance registers, training evaluations, and student drills. Manager oversees center settings, "
            "pricing matrices, promo vouchers, staff permissions, class lifecycles, and financial/capacity analytics. "
            "AI Assistant provides smart conversational guidance. Mail Delivery Service dispatches verification codes "
            "and electronic invoices. Background Scheduler executes maintenance sweeps."
        )
        print("Updated Paragraph 52 (Use Case Diagram description)")

    # 5. Update Paragraph 57 (Swimlane Diagram description)
    p57 = doc.paragraphs[57]
    if "Figure 3 is registration" in p57.text:
        p57.text = (
            "Figure 3 depicts the member registration and onboarding workflow as a five-lane UML 2.0 activity diagram: "
            "Guest, Web Client (Frontend), Arena3 API Server (Backend), Mail Delivery Service, and PostgreSQL Database. "
            "When the Guest submits registration details, the frontend validates syntax client-side and dispatches the API request. "
            "Arena3 normalizes the phone number, checks uniqueness against the database, generates a cryptographic 6-digit OTP, "
            "persists the challenge with a 5-minute TTL, and triggers asynchronous email delivery. The Guest retrieves the OTP "
            "and inputs it into the web screen; upon server-side hash verification, the active member account is created with a unique "
            "member code, the challenge record is consumed, audit logs are committed, and an authenticated JWT session is issued."
        )
        print("Updated Paragraph 57 (Swimlane Diagram description)")

    # 6. Update Table 1 (Actors)
    t1 = doc.tables[1]
    actors_data = [
        ("1", "Guest (Unregistered Visitor)", "Unregistered visitor. Browses public court schedules, tariff catalog, coach rosters, and initiates account registration with Email OTP verification."),
        ("2", "Member (Registered Customer)", "Registered active customer. Books courts, enrolls in training classes, purchases membership plans, checks in at gate via dynamic QR pass or member code, interacts with AI Assistant, and tracks personal training progress."),
        ("3", "Receptionist (Front Desk / POS)", "Front-desk personnel. Opens/closes cash till shifts, accepts cash and VietQR POS payments, scans QR passes or inputs member codes for gate check-in, rents equipment, and processes walk-ins."),
        ("4", "Coach (Sports Instructor)", "Conducts sports training classes. Views assigned teaching schedules, manages session attendance registers, records trainee skill assessments, drills, and homework."),
        ("5", "Center Manager (Operations / Admin)", "Executive administrator. Manages pricing matrices, promo vouchers, staff permissions, class lifecycles, center settings, and views/exports financial and court utilization analytics (Excel/PDF)."),
        ("6", "AI Assistant (Smart Concierge)", "In-app intelligent virtual assistant. Provides conversational guidance regarding court availability, membership plans, schedules, and complex center rules."),
        ("7", "Mail Delivery Service (Email Engine)", "External communication infrastructure. Dispatches transactional email OTP verification codes for registration and password resets, as well as e-invoices and booking alerts."),
        ("8", "Background Scheduler (Cron Engine)", "Automated background job engine. Periodically releases expired court holds, marks booking no-shows, materializes class sessions, and executes database maintenance routines."),
        ("9", "payOS Payment Gateway (VietQR)", "External banking payment gateway. Generates standard Napas247 VietQR payment links, verifies digital transaction signatures, and broadcasts signed webhook confirmations.")
    ]
    # Keep header row, adjust data rows
    while len(t1.rows) > 1:
        tr = t1.rows[-1]._tr
        t1._tbl.remove(tr)

    for item in actors_data:
        row = t1.add_row()
        row.cells[0].text = item[0]
        row.cells[1].text = item[1]
        row.cells[2].text = item[2]
    print(f"Updated Table 1 with {len(actors_data)} actors")

    # 7. Update Paragraph 111 (Section 2.2 User classes)
    p111 = doc.paragraphs[111]
    p111.text = (
        "The system defines nine actors: five primary human user classes (Guest, Member, Receptionist, Coach, Manager) "
        "and four automated/external services (AI Assistant, Mail Delivery Service, Background Scheduler, and payOS Payment Gateway)."
    )
    print("Updated Paragraph 111 (Section 2.2 User classes)")

    doc.save(docx_path)
    print("Saved modified DOCX to:", docx_path)

    # 8. Replace media images in DOCX zip package with high-resolution English diagrams
    img_replacements = {
        "word/media/image1.png": "docs/diagram-context.png",
        "word/media/image2.png": "docs/diagram-usecase.png",
        "word/media/image3.png": "docs/diagram-swimlane.png",
    }

    temp_zip = docx_path + ".tmp.zip"
    with zipfile.ZipFile(docx_path, 'r') as zin:
        with zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in img_replacements and os.path.exists(img_replacements[item.filename]):
                    print(f"Replacing {item.filename} with {img_replacements[item.filename]}")
                    with open(img_replacements[item.filename], 'rb') as f_img:
                        zout.writestr(item, f_img.read())
                else:
                    zout.writestr(item, zin.read(item.filename))

    os.replace(temp_zip, docx_path)
    print("Successfully replaced diagram images in DOCX package!")

if __name__ == "__main__":
    update_docx()
