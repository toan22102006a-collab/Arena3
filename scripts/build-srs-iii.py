# Builds docs/Arena3-SRS.docx as chapter III of the capstone report.
# Facts come from the running Arena 3 system. Figure grammar follows that chapter.
# Run: python scripts/build-srs-iii.py

import shutil
import struct
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "Arena3-SRS.docx"
BACKUP = DOCS / "Arena3-SRS-v1.4.docx"

NAVY = RGBColor(0x1B, 0x4F, 0x72)
BLACK = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CONTENT = 17.0  # A4 width 21 cm minus 2 cm margins

# role letters: G guest, M member, R receptionist, C coach, N manager
ROLE_ORDER = ("G", "M", "R", "C", "N")
ROLE_HEAD = ("Guest", "Member", "Recep.", "Coach", "Manager")


def dxa(cm):
    return str(int(round(cm * 567)))


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        edge_el = OxmlElement(f"w:{edge}")
        edge_el.set(qn("w:val"), "single")
        edge_el.set(qn("w:sz"), "4")
        edge_el.set(qn("w:space"), "0")
        edge_el.set(qn("w:color"), "CBD5E1")
        el.append(edge_el)
    tcPr.append(el)


def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def set_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), dxa(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), dxa(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), dxa(widths[min(i, len(widths) - 1)]))
            tcW.set(qn("w:type"), "dxa")


def write_cell(cell, text, bold=False, size=10, color=None, center=False, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color or BLACK)
    borders(cell)
    if fill:
        shade(cell, fill)
    margins(cell)


def add_more(cell, text, bold=False, size=10):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def header_row(row):
    trPr = row._tr.get_or_add_trPr()
    mark = OxmlElement("w:tblHeader")
    trPr.append(mark)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    mark = OxmlElement("w:cantSplit")
    trPr.append(mark)


def add_table(doc, headers, rows, widths, center_cols=()):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_widths(table, widths)
    for i, head in enumerate(headers):
        write_cell(table.rows[0].cells[i], head, bold=True, size=10, color=WHITE, center=True, fill="1B4F72")
    header_row(table.rows[0])
    cant_split(table.rows[0])
    for r, row in enumerate(rows):
        fill = "F8FAFC" if r % 2 else "FFFFFF"
        for c, val in enumerate(row):
            write_cell(
                table.rows[r + 1].cells[c],
                val,
                size=10,
                center=c in center_cols,
                fill=fill,
            )
        cant_split(table.rows[r + 1])
    spacer(doc, 6)
    return table


def spacer(doc, pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.line_spacing = 1.0
    return p


def para(doc, text, size=12, bold=False, italic=False, center=False, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def rich(doc, parts, space_after=8, space_before=0):
    """parts: list of (text, bold)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size=12, bold=bold)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def caption(doc, text):
    p = para(doc, text, size=10, italic=True, center=True, space_after=12, space_before=2)
    return p


def png_size(path):
    data = path.read_bytes()[:24]
    return struct.unpack(">II", data[16:24])


def figure(doc, path, text, max_h_cm=18.5):
    path = Path(path)
    if not path.exists():
        raise SystemExit(f"missing figure: {path}")
    w_px, h_px = png_size(path)
    width = CONTENT
    height = width * h_px / w_px
    if height > max_h_cm:
        height = max_h_cm
        width = height * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    # One dimension only: Word keeps the aspect ratio.
    run.add_picture(str(path), width=Cm(width))
    caption(doc, text)


def configure(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Times New Roman")
    sizes = {1: 16, 2: 14, 3: 13, 4: 12}
    before = {1: 16, 2: 12, 3: 10, 4: 8}
    for level, size in sizes.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before[level])
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Times New Roman")


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 13, 4: 12}[level], bold=True, color=NAVY)
    return p


def page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=NAVY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT), WD_TAB_ALIGNMENT.RIGHT)
    left = p.add_run("Arena 3  ·  Software Requirements Specification")
    set_run_font(left, size=9, color=NAVY)
    tab = p.add_run("\t")
    set_run_font(tab, size=9)
    page_field(p)
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Chapter III")
    set_run_font(hr, size=9, italic=True, color=NAVY)
    # bottom border on the header paragraph
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1B4F72")
    pBdr.append(bottom)
    jc = pPr.find(qn("w:jc"))
    if jc is not None:
        jc.addprevious(pBdr)
    else:
        pPr.append(pBdr)


def toc(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Open the document in Word and update this field to refresh the contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(begin)
    r.append(instr)
    r.append(separate)
    r.append(text)
    r.append(end)
    set_run_font(run, size=11, italic=True)


# --- content -----------------------------------------------------------------

ACTORS = [
    ("1", "Guest", "A person with no session. They read the public schedule and register. They do not book, pay, or see another person's data."),
    ("2", "Member", "A registered customer, with or without a current plan. They book courts, join classes, buy and freeze plans, and read their own receipts."),
    ("3", "Receptionist", "Front-desk staff. They open a till, take money, check the gate, sell walk-ins and gear, and reconcile transfers."),
    ("4", "Coach", "Teaches assigned classes only. They mark the register, record results, and write plans and homework. They do not take money."),
    ("5", "Manager", "Runs the centre. They set prices, plans, staff, and flags, sign off large refunds, and read reports and the audit log. They can also open desk and coach screens."),
]

# id, name, actors, description, trigger, flow, rules
USE_CASES = [
    ("UC-01", "Register member account", "Guest",
     "Create a member account. The account row is written only after UC-02 accepts the code.",
     "The guest opens Register.",
     "The guest enters name, phone, date of birth, and a password typed twice, and accepts the terms and the privacy notice. A person under minor_age (default 16) also enters a guardian name and phone. The system stores an OTP challenge, not the account.",
     "BR-01 one phone, one account. BR-02 password of at least 8 characters with letters and digits. BR-07 guardian for a minor. BR-08 terms and privacy must be accepted."),
    ("UC-02", "Verify email OTP code", "Guest",
     "Included by UC-01. Confirms the email challenge and then creates the member.",
     "Registration has just issued a challenge.",
     "The guest types the code. A match creates the member and opens Member Home. The challenge lasts 5 minutes. Five wrong attempts lock it for 15 minutes.",
     "There is no SMS channel. In production the code is not printed on the screen. The email outbox transport is still a stub, so a deployment with real customers must attach a real transport before relying on delivery."),
    ("UC-03", "Log in and reset password", "Guest, Member, Receptionist, Coach, Manager",
     "Open a session, or replace a forgotten password. The figure draws Guest and Member beside this oval; every role uses the same screen.",
     "Someone opens Login or Forgot password.",
     "Sign-in takes the phone and password and opens that role's home. A staff account that must change its password still lands on the role home; Account opens on the Security tab while the flag is set. Forgot-password answers the same whether or not the phone exists, then accepts a code and a new password.",
     "Five wrong passwords lock the account for 15 minutes. A member session lasts 7 days. A staff session lasts 12 hours. Passwords are stored as salted hashes."),
    ("UC-04", "Display the QR member pass", "Member, Receptionist",
     "Included by gate check-in. The member shows the pass and the desk reads it.",
     "Check-in needs to identify the member.",
     "The pass is shown from the member side. The desk can also look the member up by phone or member code.",
     "BR-54 a gate check-in never writes class attendance."),
    ("UC-05", "Browse the court schedule", "Guest, Member",
     "Read which courts are free. Choosing a slot is UC-06 and requires a member session.",
     "A visitor or member opens the schedule.",
     "The schedule shows free, held, and busy slots, plus a week strip and a month grid of up to 62 days. Days in the past, or beyond book_ahead_days (default 7), are visible and cannot be selected.",
     "A member never receives the staff occupancy detail (who, phone, amount)."),
    ("UC-06", "Book a court and hold the slot", "Member",
     "Reserve one court-time. Includes UC-13 when the member pays online. A reschedule is the extension UC-09.",
     "A member chooses a free slot.",
     "The system creates a hold, not a confirmed booking. The default hold is 5 minutes (hold_minutes). Paying, or spending a plan hour, confirms it in the same transaction as the payment and the invoice.",
     "BR-32 horizon and at most max_slots_per_day (default 2) slots that day. BR-35 inside opening hours, court ready. BR-66 no past slot. BR-44 debt over debt_limit_vnd blocks a new booking. BR-17 a plan hour may be spent instead of money. A second member who wants the same court-time receives 409 CONFLICT_SLOT."),
    ("UC-07", "Cancel a court booking", "Member, Receptionist, Manager",
     "Release a hold or a confirmed booking. Refund, when one is due, is the extension UC-08.",
     "The holder, or the desk, cancels a hold or a confirmed booking.",
     "A hold is released with no refund. A confirmed booking is priced with the cancellation tiers and the occupancy is released. Reception or a manager may waive the fee.",
     "Default tiers: 100 percent when at least 24 hours remain, 50 percent when at least 4 and under 24 hours remain, none under 4 hours. The 2-hour cancel_court_hours window blocks a move (UC-09), not a cancellation. A no-show is not refunded."),
    ("UC-08", "Refund the court booking", "Member, Receptionist, Manager",
     "Extends UC-07. Pays back the tier amount, or returns a plan hour when the tier is at least 50 percent.",
     "A cancellation of a paid booking produces a refund amount above zero, or the desk raises a refund on a payment.",
     "A refund at or above refund_manager_vnd (default 1,000,000 VND) waits for a manager. Below that threshold it is posted immediately. The member sees the payment as refunded.",
     "BR-style guard FR-D22: the same payment cannot be refunded twice. The guard sums the ledger for that booking. A price-changing reschedule is not a refund; it is refused, and the booking stays on the old slot."),
    ("UC-09", "Reschedule the booking", "Member",
     "Extends UC-06. Move a confirmed booking to another slot without a second charge.",
     "The member chooses a new slot for a booking they already hold.",
     "The move and the occupancy change commit together. A lost race returns 409 and leaves the booking where it was. A different new price is refused and both amounts are shown; the member cancels and books again. A booking paid with a plan hour has no price difference.",
     "The move is refused inside cancel_court_hours (default 2) before the start. The new slot still has to satisfy BR-32, BR-35, and BR-66."),
    ("UC-10", "View the court matrix", "Receptionist, Manager",
     "See every court and who holds it, including a walk-in.",
     "The desk opens the court map.",
     "Each cell opens the holder, the code, the time, the amount, and the status. A free convertible pair can be merged when both halves are free.",
     "Members do not have this screen."),
    ("UC-11", "Buy or renew a membership", "Member",
     "Order a plan. Includes UC-13 when the order is paid online. The desk can sell the same plan inside UC-22.",
     "A member chooses a plan that is on sale.",
     "The order creates a pending subscription. It becomes active when enough has been paid, in full or up to the deposit the centre allows. Withdrawn plans cannot be ordered.",
     "BR-65 a withdrawn plan stays in history and is not hard-deleted. BR-14 a frozen plan blocks a new purchase until it is resumed. Orders are limited to 3 a minute and 10 an hour."),
    ("UC-12", "Freeze or resume a membership", "Member",
     "Pause an active plan, then resume it. Days used count against the yearly allowance.",
     "A member freezes or resumes their active plan.",
     "Only an active plan can be frozen. While it is frozen the member cannot buy another plan.",
     "BR-14 and BR-15. The yearly cap is freeze_max_days_year."),
    ("UC-13", "Pay online by VietQR", "Member, Receptionist",
     "Included by booking and by buying a plan. Collects money through payOS. The return page is not proof of payment.",
     "The member, or the desk, asks for a payment link on a held court or a pending plan.",
     "Arena 3 creates one live link and shows the QR. The payment is posted only after a poll or a signed webhook says the amount paid is at least the amount due. Settlement then follows the same path as cash: booking confirmed or plan activated, invoice issued, payer notified.",
     "If payOS is not configured the request is refused and the desk takes the money. A second press returns the same link. An expired or cancelled link does not post money. The amount on the return URL is never trusted by itself."),
    ("UC-14", "Enroll in a training class", "Member",
     "Join an open class, or take a numbered place on its waitlist.",
     "A member enrols, accepts a waitlist offer, or leaves a class.",
     "A full class offers the next waitlisted member the seat for waitlist_offer_hours (default 2). An expired offer moves on. Leaving is allowed until cancel_class_hours (default 4) before the next session.",
     "BR-12 the plan must cover the sport and be active. BR-18 session quota remaining. BR-24 no clash and no second enrolment. BR-67 the class is open. BR-22 do not overfill, including between offer and acceptance. BR-25 an expired offer cannot be accepted. BR-20 the leave window. BR-26 if the centre cancels, a per-session plan receives that session back once."),
    ("UC-15", "View the assigned coach schedule", "Coach",
     "Read this coach's sessions, grouped by day.",
     "The coach opens their home.",
     "Each session shows date, time, court, class code, enrolment, and level. Opening one opens its register.",
     "BR-59 a coach sees only classes they lead or assist."),
    ("UC-16", "Record session attendance", "Coach",
     "Mark each enrolled student for one session. Includes UC-17 when the register is submitted.",
     "The coach opens a session they teach.",
     "Each student is Present, Late, Absent, or Excused. The member can later read only their own marks. A session not yet marked reads as not marked.",
     "BR-59 own classes only. BR-57 Late is the coach's mark. BR-58 three Absents in a row raise an alert; Excused breaks the streak; enrolment is not dropped. BR-54 the gate does not fill this register."),
    ("UC-17", "Lock the attendance book", "Coach, Manager",
     "Included by UC-16, and also closed by the scheduled job when the session has been over for two hours.",
     "The coach submits the register, or two hours pass after the session end.",
     "A locked register refuses the coach. A manager may correct it.",
     "BR-53 the reason for a correction is at least 3 characters and is audited."),
    ("UC-18", "Assess a trainee", "Coach",
     "Record results, a session plan, homework, and a review. Nothing here takes money.",
     "The coach opens a student from the register, or writes the session panels.",
     "The member sees level, recent results, reviews, and open homework on Progress. Coach notes stay on the staff profile.",
     "BR-56 training does not charge and does not change a plan balance. BR-55 health notes stay inside the viewer's scope. BR-62 when flag F4 is off, every training route answers 403."),
    ("UC-19", "Manage class sessions", "Manager",
     "Publish a class and change its life: one session, the coach, the capacity, or the whole class.",
     "A manager creates, edits, moves, or cancels a class or a session.",
     "Creating a class picks a real coach and a real court. Cancelling a session or a class returns one session per enrolled member on a per-session plan, once, and tells members and the waitlist. A move keeps the original time on record.",
     "BR-21 no coach or court clash. BR-23 the coach teaches that sport. BR-27 a substitute takes the class from now on and must be free. BR-28 a move inside 12 hours needs a reason. BR-22 capacity cannot drop below the number enrolled. BR-26 the session is returned once. BR-67 closed classes take no new enrolment."),
    ("UC-20", "Run a till shift", "Receptionist, Manager",
     "Open the drawer before taking money, and close it against a cash count.",
     "The receptionist starts a shift, or closes the open one.",
     "Every payment they take is attached to that shift. Closing stores the counted cash next to the expected total.",
     "BR-49 no shift, no money. Only one shift per receptionist is open at a time."),
    ("UC-21", "Check in at the gate", "Receptionist",
     "Record that a member entered the building. Includes UC-04.",
     "The desk checks a member in by phone or member code.",
     "The visit is stored once per five minutes and listed for today.",
     "BR-54 this is not class attendance. BR-62 the gate route requires flag F4."),
    ("UC-22", "Sell at the front desk", "Receptionist",
     "Find or create a member, sell a plan at the counter, and answer a support request.",
     "A person is at the counter, or a request is open on the desk home.",
     "New-member selling is one flow: profile, plan, payment. A reply closes the request and notifies the member. A duplicate can be closed with no reply. The reply names the staff member who wrote it.",
     "A phone already used by another account is refused (BR-01)."),
    ("UC-23", "Rent gear and walk-ins", "Receptionist",
     "Sell a court to someone with no account, and lend equipment.",
     "A walk-in wants a court, or someone borrows gear.",
     "A walk-in is name and phone only. Gear goes to a member found by search, or to a guest phone. Stock falls on the way out and returns by the quantity on that loan.",
     "BR-66 a walk-in into a slot already running needs at least 20 minutes left. BR-38 stock cannot go negative. Returning the same loan twice is 409. Gear rental does not write a payment row, so it is absent from revenue."),
    ("UC-24", "Reconcile bank transfers", "Receptionist, Manager",
     "Turn a transfer the member claims to have sent into a real payment, or reject it.",
     "A member chose bank transfer, or the desk opens the transfer queue.",
     "Choosing transfer posts no payment and confirms nothing (HTTP 202). The hold lasts transfer_hold_minutes (default 120). The queue is oldest first and shows the phone. Reconciling posts the same payment, invoice, and notice as any other method. Rejecting returns the slot and tells the member.",
     "A member cannot reconcile their own transfer. A hold that has already expired cannot be reconciled."),
    ("UC-25", "Set prices and vouchers", "Manager",
     "Maintain price rules, plans on sale, opening rules, courts, flags, and promotion codes.",
     "A manager edits prices, a plan, settings, a court, or a promotion.",
     "A price rule is keyed by sport, court class, day type, and hour band. A promotion is a percent or a fixed amount. A percent is applied once and rounded once to the nearest 1,000 VND, and the amount due never falls below 1,000 VND.",
     "BR-34B no matching price rule means the quote is refused; a default price is not invented. BR-46 the discount rounding above. BR-65 stop-selling withdraws a plan. Time zone and currency are not editable: one centre, Asia/Ho_Chi_Minh, VND."),
    ("UC-26", "Manage staff accounts", "Manager",
     "Issue receptionist and coach accounts. There is no public staff sign-up and no separate admin role.",
     "A manager creates, locks, unlocks, re-roles, or resets a staff account.",
     "The first password is temporary. The hand-over screen shows it once. Resetting a password forces a change at the next sign-in and can revoke sessions.",
     "A manager cannot lock or demote themselves. Every action is audited."),
    ("UC-27", "View revenue and utilization", "Manager",
     "Read revenue, capacity, and members for a chosen period, next to the previous period of the same length. Includes UC-28.",
     "A manager opens Reports.",
     "Revenue splits by source and by till, and can be filtered by method. A rise in revenue is favourable; a rise in refunds is not. Capacity is a heatmap of up to 93 days. Attendance percent is hidden when nothing has been marked, or when F4 is off, rather than shown as zero.",
     "Court revenue is kept apart from class fees. Class fees and gear currently contribute no payment row, so they show as zero."),
    ("UC-28", "Export a report", "Manager",
     "Included by UC-27. Download the same figures the screen shows.",
     "The manager chooses Excel, PDF, or CSV for the current report.",
     "The file is built on request from the same queries as the screen and is not stored. Spreadsheet money cells are numbers.",
     "An unknown format is 400. An unknown report is 404. Receptionists and coaches are refused. The PDF embeds a font so Vietnamese diacritics render."),
    ("UC-29", "Read the audit log", "Manager",
     "Read who changed a sensitive record. Rows are never deleted from this screen.",
     "A manager opens Audit and filters by actor, action, entity, or dates.",
     "The default list is short, with a control to show more. Hiding a row is only a view.",
     "Staff state changes write an audit row: refunds, transfer decisions, shift close, settings, ticket replies, session and class changes, staff accounts, attendance corrections, and online-payment links."),
    ("UC-30", "Ask the AI assistant", "Member",
     "Ask about hours, prices, plans, coaches, classes, booking, cancel, and waitlists. The assistant cannot change anything.",
     "A member opens Assistant and sends a question.",
     "The answer is built from the centre's own settings, plans, classes, and coaches. No member's personal data is sent to the model. Without a model key, or if the call fails, the assistant answers from local rules and labels the source. It can open a support request for the desk.",
     "A question is at most 800 characters. At most 20 questions a minute. No medical advice. The assistant has no tool that takes money, books, cancels, or edits a record."),
    ("UC-31", "Run scheduled jobs", "Manager",
     "The figure associates this goal with Manager because the manager owns the operating rules. No person signs in to run a pass. The ten jobs are specified in section 3.1.4.",
     "The process wakes on its timer.",
     "A pass first asks, in one read, whether any job has due rows. Only those jobs then run, each in its own transaction.",
     "The loop is about every 15 seconds. Generating class sessions is further limited to every 10 minutes."),
]

FLOWS = [
    ("3.1.1.1 Public and authentication", "docs/drawio/01-Public-Auth.png",
     "Figure 2. Screen flow for a visitor: landing, register, sign-in, and password reset, then the home of the signed-in role. The payment-return page is listed in section 3.1.2.1; it is opened by payOS, not by this flow. Source: docs/drawio/01-Public-Auth.drawio."),
    ("3.1.1.2 Member", "docs/drawio/02-Member.png",
     "Figure 3. Screen flow after a member signs in. Home is /app. Source: docs/drawio/02-Member.drawio."),
    ("3.1.1.3 Receptionist", "docs/drawio/03-Receptionist-Desk.png",
     "Figure 4. Screen flow for the front desk. Home is /desk. Source: docs/drawio/03-Receptionist-Desk.drawio."),
    ("3.1.1.4 Coach", "docs/drawio/04-Coach.png",
     "Figure 5. Screen flow for a coach. Home is /coach. Source: docs/drawio/04-Coach.drawio."),
    ("3.1.1.5 Manager", "docs/drawio/05-Manager.png",
     "Figure 6. Screen flow for a manager. Home is /manager. Desk screens a manager opens from this menu are Member Profile and Payments. The other desk and coach screens are not drawn again here. Source: docs/drawio/05-Manager.drawio."),
]

# group heading, list of (feature, screen, description, role letters)
SCREENS = [
    ("3.1.2.1 Public and authentication", [
        ("Landing", "Landing Page (/)", "Public home of the centre. States what the centre offers and links to Login and Register.", "G"),
        ("Landing", "Coach Profile Modal", "Opens from the coach list on the landing page and shows one coach.", "G"),
        ("Login", "Login Screen (/login)", "Phone and password. Success opens that role's home. A staff password that must be changed does not divert the login; Account opens on Security while the flag is set.", "GMRCN"),
        ("Register", "Registration Form (/register)", "A new member enters name, phone, date of birth, and a password typed twice.", "G"),
        ("Register", "OTP Verification Step", "The one-time code. The account is created only after a match. Five minutes, then five wrong attempts lock the challenge for 15 minutes.", "G"),
        ("Forgot password", "Request Code Step (/forgot)", "The person enters a phone. The reply does not reveal whether an account exists.", "GMRCN"),
        ("Forgot password", "Code and New Password Step", "Code plus a new password, then back to Login.", "GMRCN"),
        ("Online payment", "Payment Return Screen (/pay/return)", "Shown when the browser comes back from payOS. It confirms nothing by itself and returns the person to the screen they left.", "M"),
    ]),
    ("3.1.2.2 Member", [
        ("Home", "Member Home (/app)", "Upcoming bookings, the active plan, and shortcuts.", "M"),
        ("Home", "Waitlist Offer Card", "A seat opened on a waited class. Accept or decline.", "M"),
        ("Payment", "Scan-to-Pay Modal", "QR and amount for an outstanding payment. Closes when the payment is confirmed.", "M"),
        ("Court booking", "Book Court (/app/book)", "Date, court type, and slot, then payment to confirm.", "M"),
        ("Court booking", "Month Calendar View", "Which days still have space.", "M"),
        ("Court booking", "Hold Panel", "Holds the slot for the configured minutes (the member screen shows 5) so nobody else can take it.", "M"),
        ("Court booking", "Clash Confirmation", "Warns when the new booking overlaps one of the member's own bookings.", "M"),
        ("Court booking", "Alternative Slots", "Nearby free slots when the chosen one is gone.", "M"),
        ("Court booking", "Change Time Mode", "Moves an existing booking instead of creating one.", "M"),
        ("Court booking", "Scan-to-Pay Modal", "QR for this booking. The browser then hits the payment-return screen.", "M"),
        ("Classes", "Classes (/app/classes)", "Join, wait, or leave a coached class.", "M"),
        ("Classes", "Waitlist Offer Card", "The same offer, beside the class it belongs to.", "M"),
        ("Classes", "My Attendance Panel", "The member's own marks. Nobody else's.", "M"),
        ("Plans", "Plans (/app/plans)", "The current plan and plans on sale.", "M"),
        ("Training", "Progress (/app/train)", "Level, session results, and homework. Hidden when F4 is off.", "M"),
        ("Assistant", "Assistant (/app/assistant)", "Questions answered from the centre's own timetable, plans, and coaches.", "M"),
        ("Notifications", "Notifications (/app/notifications)", "The member's notices.", "M"),
        ("Notifications", "Notification Detail Modal", "Full text and a link to the related booking, class, plan, progress, receipt, or support thread.", "M"),
        ("Account", "Account Settings (/account)", "Shared by every signed-in role.", "MRCN"),
        ("Account", "Profile Tab", "Own name and health notes.", "MRCN"),
        ("Account", "Security Tab", "Change password. Shown first while a temporary password is outstanding.", "MRCN"),
        ("Account", "Receipts Tab", "Every invoice. Each opens as a PDF.", "M"),
        ("Account", "Support Tab", "Write to the desk and read replies.", "M"),
        ("Receipts", "Receipt PDF", "Printable receipt, opened in a new tab.", "M"),
    ]),
    ("3.1.2.3 Receptionist", [
        ("Home", "Desk Home (/desk)", "The open shift and the queues.", "RN"),
        ("Shift", "Close Shift Modal", "Counted cash against the book total.", "RN"),
        ("Onboarding", "New Member Wizard", "Three steps at the counter: profile, plan, payment.", "RN"),
        ("Support", "Support Request Reply", "Read a request, reply, or close it.", "RN"),
        ("Check-in", "Gate Check-in (/desk/gate)", "Who entered today. Does not mark a class.", "RN"),
        ("Courts", "Court Map (/desk/courts)", "Live occupancy of every court.", "RN"),
        ("Courts", "Occupancy Detail Modal", "Who holds the cell, for how long, and for how much.", "RN"),
        ("Courts", "Walk-in Payment Form", "Book and charge a person with no account.", "RN"),
        ("Courts", "Merge BR and BC", "Join a free convertible pair into one court.", "RN"),
        ("Classes", "Classes (/desk/classes)", "Read-only classes and rosters.", "RN"),
        ("Classes", "Class Detail Modal", "One class, its sessions, and who is enrolled or waiting.", "RN"),
        ("Payments", "Payments (/desk/payments)", "Transfers, waiting payments, refunds, and receipts.", "RN"),
        ("Payments", "Transfers to Check", "Bank transfers waiting for a human to match a statement.", "RN"),
        ("Payments", "Waiting for Payment", "Holds and orders not yet paid.", "RN"),
        ("Payments", "Refund Queue", "Refunds, including those waiting for a manager.", "RN"),
        ("Payments", "Receipts List", "Issued receipts, each a PDF.", "RN"),
        ("Payments", "Scan-to-Pay Modal", "A QR shown at the counter.", "RN"),
        ("Gear", "Gear (/desk/gear)", "Lend and take back. Stock moves with the loan.", "RN"),
        ("Members", "Member Profile (/desk/member/:id)", "Plan, bookings, and payments for one member.", "RN"),
        ("Members", "Edit Profile Modal", "Correct name, phone, date of birth, and guardian.", "RN"),
        ("Members", "Raise Refund Modal", "Start a refund on one payment.", "RN"),
        ("Alerts", "Alerts (/alerts)", "Shared by every signed-in role.", "MRCN"),
        ("Account", "Account Settings (/account)", "Own profile and password.", "MRCN"),
    ]),
    ("3.1.2.4 Coach", [
        ("Schedule", "Coach Schedule (/coach)", "This coach's sessions. Home of the role.", "CN"),
        ("Attendance", "Attendance (/coach/attendance)", "One session's register, results, plan, and homework.", "CN"),
        ("Attendance", "Session Register", "Present, Late, Absent, or Excused. Locks on submit.", "CN"),
        ("Attendance", "Results Panel", "Session results for the students.", "CN"),
        ("Attendance", "Session Plan Panel", "The plan for this session.", "CN"),
        ("Attendance", "Homework Panel", "Homework for the class or for one student.", "CN"),
        ("Students", "Student Profile (/coach/student/:id)", "Training record opened from the register.", "CN"),
        ("Students", "Level", "Current level per sport.", "CN"),
        ("Students", "Progress Review", "Reviews over time.", "CN"),
        ("Students", "Coach Notes", "Notes the member never receives.", "CN"),
        ("Students", "Homework History", "Assigned work and whether it is done.", "CN"),
        ("Alerts", "Alerts (/alerts)", "Includes a repeated-absence alert.", "CN"),
        ("Account", "Account Settings (/account)", "Own profile and password.", "CN"),
    ]),
    ("3.1.2.5 Manager", [
        ("Reports", "Reports (/manager)", "Home of the manager. A period against the previous period.", "N"),
        ("Reports", "Revenue and Occupancy Summary", "Totals for money and how full the courts were.", "N"),
        ("Reports", "Capacity Heatmap", "Fullness by day and hour.", "N"),
        ("Reports", "Members and Enrolment", "New members, renewals, enrolment, attendance percent.", "N"),
        ("Reports", "Export", "Excel, PDF, or CSV of the report on screen.", "N"),
        ("Classes", "Classes (/manager/classes)", "Create and manage classes.", "N"),
        ("Classes", "Create Class Form", "Coach and court chosen from real rows.", "N"),
        ("Classes", "Class Detail Modal", "Holds the four forms below.", "N"),
        ("Classes", "Cancel or Move Session", "One session, with a reason when the move is inside 12 hours.", "N"),
        ("Classes", "Change Coach", "Another coach from now on, free and qualified.", "N"),
        ("Classes", "Edit Capacity", "Cannot drop below the number already enrolled.", "N"),
        ("Classes", "Cancel Class", "Stops the class, returns one session each, notifies roster and waitlist.", "N"),
        ("Members", "Members (/manager/members)", "Searchable directory with plan, debt, and class count.", "N"),
        ("Members", "Member Profile", "The same desk profile, opened from the directory.", "N"),
        ("Payments", "Payments", "The desk payments screen, used here to sign off refunds.", "N"),
        ("Plans", "Plans (/manager/plans)", "Plans the centre sells.", "N"),
        ("Plans", "New Plan Modal", "Create a plan.", "N"),
        ("Plans", "Plan Detail Modal", "Price, duration, hours or sessions, discount, sport, and how many members hold it.", "N"),
        ("Staff", "Staff (/manager/staff)", "Receptionist and coach accounts.", "N"),
        ("Staff", "Add Staff Modal", "Create an account.", "N"),
        ("Staff", "Role Modal", "Change role between receptionist and coach.", "N"),
        ("Staff", "Hand-over Password Modal", "Shows the temporary password once.", "N"),
        ("Pricing", "Pricing (/manager/prices)", "Court and service price rules.", "N"),
        ("Audit", "Audit (/manager/audit)", "Sensitive staff actions. Nothing here deletes a row.", "N"),
        ("Settings", "Settings (/manager/settings)", "Hours, windows, limits, courts, and feature switches.", "N"),
        ("Alerts", "Alerts (/alerts)", "Manager notices.", "N"),
        ("Account", "Account Settings (/account)", "Own profile and password.", "N"),
    ]),
]

JOBS = [
    ("1", "Bookings", "Expire holds", "A booking in status hold whose hold_until is in the past is cancelled and its occupancy is released. The slot goes back on sale. No refund, because a hold was never paid."),
    ("2", "Bookings", "Mark no-shows", "A confirmed booking still not in use after start plus noshow_grace_minutes becomes no_show and releases the court. A no-show is not refunded."),
    ("3", "Bookings", "Complete bookings", "A booking in use whose end was more than 10 minutes ago is completed."),
    ("4", "Membership", "Expire subscriptions", "An active subscription whose end_on is before today's date in Asia/Ho_Chi_Minh becomes expired."),
    ("5", "Membership", "Expiry reminders", "Writes an in-app notice at 7, 3, and 0 days before the end. A dedupe key stops the same notice being queued twice."),
    ("6", "Notices", "Flush the outbox", "Sends queued notices whose attempt count is under 5, on the in-app channel and the email channel. sent_at is set only after a successful attempt. The email transport is a stub; the in-app channel is the one members actually see."),
    ("7", "Payments", "Expire payOS links", "Runs only when payOS is configured. Cancels a link whose court hold has already lapsed."),
    ("8", "Classes", "Expire waitlist offers", "A pending offer past expires_at is expired, that waitlisted enrolment is cancelled, and the next person is invited."),
    ("9", "Classes", "Lock attendance", "A scheduled session whose end was more than 2 hours ago is closed through the same release used when a register is finished."),
    ("10", "Classes", "Generate sessions", "For every open class, materialise upcoming sessions from its rule. Limited to once every 10 minutes, inside the shorter job loop."),
]

RULES = [
    ("BR-01", "One account per phone number."),
    ("BR-02", "Password of at least 8 characters, with letters and digits, typed twice when it is set."),
    ("BR-07", "A person under minor_age (default 16) needs a guardian name and phone."),
    ("BR-08", "Terms and the privacy notice must be accepted at registration."),
    ("BR-12", "Class enrolment needs an active plan that covers that sport."),
    ("BR-14", "Only an active plan can be frozen, and a frozen plan blocks a new purchase."),
    ("BR-15", "Freeze days count against freeze_max_days_year."),
    ("BR-17", "Paying with a plan requires court hours left. One hour is consumed."),
    ("BR-18", "A session-quota plan must have sessions left."),
    ("BR-20", "A member may leave a class until cancel_class_hours (default 4) before the next session."),
    ("BR-21", "A court or a coach cannot be in two places at the same time."),
    ("BR-22", "A class cannot be filled past capacity, and capacity cannot be set below the number enrolled."),
    ("BR-23", "A coach can be assigned only in a sport they teach."),
    ("BR-24", "No clashing enrolment and no second enrolment in the same class."),
    ("BR-25", "A waitlist offer expires and then cannot be accepted."),
    ("BR-26", "When the centre cancels a session or a class, a per-session plan receives that session back once."),
    ("BR-27", "A class can be handed to another coach from now on."),
    ("BR-28", "Moving a session less than 12 hours before its old time needs a reason."),
    ("BR-32", "Bookings stay inside book_ahead_days (default 7) and max_slots_per_day (default 2)."),
    ("BR-34B", "No matching price rule means the quote is refused. A default price is not invented."),
    ("BR-35", "A slot must fall inside opening hours, and the court must be ready."),
    ("BR-38", "Equipment cannot be lent beyond stock."),
    ("BR-39", "A confirmed, held, or in-use booking has an occupancy. A completed, cancelled, or no-show booking does not."),
    ("BR-44", "Debt above debt_limit_vnd blocks a new booking."),
    ("BR-46", "A percent promotion is applied once and rounded once to the nearest 1,000 VND. At least 1,000 VND remains due."),
    ("BR-49", "Money at the desk requires an open till shift."),
    ("BR-53", "The register closes 2 hours after the session ends, or when it is submitted. Only a manager, with a reason of at least 3 characters, can correct it."),
    ("BR-54", "Gate check-in does not write class attendance."),
    ("BR-55", "Health notes are visible only inside the viewer's scope."),
    ("BR-56", "Training records do not charge money and do not change a plan balance."),
    ("BR-57", "Late is a mark the coach writes on the register."),
    ("BR-58", "Three Absents in a row raise an alert. Excused breaks the streak. Enrolment is not dropped automatically."),
    ("BR-59", "A coach works only classes they lead or assist."),
    ("BR-62", "A switched-off feature flag answers 403 with br BR-62."),
    ("BR-65", "A withdrawn plan cannot be ordered and cannot be hard-deleted."),
    ("BR-66", "No booking in the past. A walk-in into a slot already running needs at least 20 minutes left."),
    ("BR-67", "A class must be open to take a new enrolment."),
]

MESSAGES = [
    ("MSG-01", "UNAUTHENTICATED", "401", "No session, or the session has expired."),
    ("MSG-02", "VALIDATION", "400", "A form value is malformed. The body names the field so the screen can mark it."),
    ("MSG-03", "FORBIDDEN", "403", "Wrong role, or another person's data. With br BR-62, the module is switched off."),
    ("MSG-04", "NOT_FOUND", "404", "No such record."),
    ("MSG-05", "CONFLICT_SLOT", "409", "Someone else already holds that court-time, or the coach is busy."),
    ("MSG-06", "HOLD_EXPIRED", "409", "The hold clock ran out. The court may already have been sold."),
    ("MSG-07", "CONFLICT_STATE", "409", "The record exists but is not in a state that accepts this action."),
    ("MSG-08", "BR_VIOLATION", "422", "A named business rule refused the action. The body carries br."),
    ("MSG-09", "VALIDATION", "422", "The request is well formed and still not acceptable in the current state."),
    ("MSG-10", "RATE_LIMITED", "429", "Too many attempts. Used for plan orders and for the assistant."),
    ("MSG-11", "Accepted, not paid", "202", "A bank transfer was chosen. No payment was posted and no booking was confirmed."),
]


def lab(doc, label, text):
    rich(doc, [(label + ". ", True), (text, False)], space_after=4)


def feature(doc, number, title, trigger, actors, purpose, interface, collect, compute, binding, layout, details):
    heading(doc, f"{number} {title}", 2)
    lab(doc, "Function trigger", trigger)
    para(doc, "Function description", bold=True, space_after=2)
    lab(doc, "Actors", actors)
    lab(doc, "Purpose", purpose)
    lab(doc, "Interface", interface)
    para(doc, "Data processing", bold=True, space_after=2)
    lab(doc, "Data collect", collect)
    lab(doc, "Data computation", compute)
    lab(doc, "Data binding", binding)
    lab(doc, "Screen layout", layout)
    lab(doc, "Function details", details)


def marks(letters):
    return ["x" if role in letters else "" for role in ROLE_ORDER]


def build():
    if OUT.exists() and not BACKUP.exists():
        shutil.copy2(OUT, BACKUP)

    doc = Document()
    configure(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    footer(section)

    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    props = doc.core_properties
    props.title = "Arena 3 Software Requirements Specification"
    props.subject = "Chapter III — Software Requirement Specification"
    props.category = "SRS"
    props.comments = "Version 2.0. Arranged as chapter III. Does not reuse the previous 60-page grade."

    # Cover
    para(doc, "FPT UNIVERSITY", size=12, bold=True, center=True, space_after=2)
    para(doc, "CAPSTONE PROJECT REPORT", size=12, center=True, space_after=18)
    para(doc, "CHAPTER III", size=14, bold=True, center=True, space_before=24, space_after=6)
    para(doc, "SOFTWARE REQUIREMENT SPECIFICATION", size=20, bold=True, center=True, space_after=10)
    para(doc, "Arena 3 Sports Complex", size=16, center=True, space_after=4)
    para(doc, "One schedule and one ledger for court hire, classes, memberships, and the front desk.", size=12, italic=True, center=True, space_after=18)
    add_table(doc,
              ["Item", "Value"],
              [
                  ["Document", "Software Requirements Specification, chapter III only"],
                  ["System", "Arena 3 sports-complex management"],
                  ["Version", "2.0"],
                  ["Date", "11 October 2026"],
                  ["Status", "As built. Describes the running system, not a future wish list."],
                  ["Supersedes", "The chapter layout of Arena3-SRS.docx version 1.4 (UC-01 through UC-37)."],
              ],
              [4.0, 13.0])
    para(doc, "This file is the requirements chapter. It is not the whole capstone report, and it does not include the introduction, the design chapter, or the test chapter. Class, sequence, and activity diagrams remain design models and are named once, at the end of section 3.1.5.", size=11, italic=True)

    doc.add_page_break()
    heading(doc, "Document history", 1)
    add_table(doc,
              ["Version", "Date", "What changed"],
              [
                  ["1.1", "Earlier narrative", "docs/SRS.md. Left unchanged on purpose. It is the long-form source this chapter was checked against."],
                  ["1.4", "5 October 2026", "Previous Word file. Numbered UC-01 through UC-37 and used an older figure set."],
                  ["2.0", "11 October 2026", "Rebuilt to the chapter III outline. Thirty-one use cases, matching the current use-case diagram. Screen flows taken from the five role diagrams. Business rules aligned with the code where version 1.1 had drifted."],
              ],
              [2.4, 3.6, 11.0])
    para(doc, "Two corrections against the version 1.1 narrative are deliberate. A member may cancel a confirmed court booking at any time before it is finished; the refund tier depends on how much notice is left. The setting cancel_court_hours (default 2) refuses a reschedule, not a cancellation. Sign-in always opens the role home. An account that must change its password is not diverted to Account during login; the Security tab is selected when Account is opened.")

    heading(doc, "Contents", 1)
    toc(doc)
    doc.add_page_break()

    # 1
    heading(doc, "1. Product Overview", 1)
    para(doc, "Arena 3 is the operating system of a single sports centre. The centre sells four things that share one set of courts: casual court hire, coached classes, membership plans, and whatever the front desk takes money for. On paper those four overwrite each other. A court is sold twice, a class is taught on a court a member has booked, a payment has nothing the member can show, and the month ends with no trustworthy total. Arena 3 keeps one schedule and one ledger, so a court-time has one holder and a payment has one invoice.")
    para(doc, "The product is a web application. After sign-in each role lands on its own home: a visitor on the public page, a member on /app, a receptionist on /desk, a coach on /coach, and a manager on /manager. Account and Alerts are shared by everyone who has a session. Money is whole dong. Dates that matter to the business are dates in Asia/Ho_Chi_Minh. The phone number is the login identity.")
    para(doc, "In scope:", bold=True, space_after=2)
    for line in [
        "Court availability, short holds, confirmed bookings, check-in, reschedule, cancellation, and refunds.",
        "Classes: publishing a timetable, enrolment, waitlists, attendance, and the class life cycle.",
        "Training, behind feature flag F4: results, plans, homework, levels, reviews, absence alerts, and the gate.",
        "Memberships: plans, orders, activation, freeze, and quota hours or sessions.",
        "The front desk: till shifts, cash, bank transfer, payOS QR, refunds, receipts, walk-ins, and equipment.",
        "Management: prices, promotions, opening rules, staff accounts, the member directory, reports, export, and the audit log.",
        "A member writing to the desk, and an assistant that answers from the centre's own data and cannot change it.",
    ]:
        bullet(doc, line)
    para(doc, "Out of scope:", bold=True, space_before=8, space_after=2)
    for line in [
        "More than one centre. Settings are a single row.",
        "Card acquiring and any bank API other than payOS. A manual transfer is still matched by a person.",
        "Outbound SMS. That channel was removed. The email side of the outbox is a stub; in-app notices are what members see today. Registration still uses an email one-time code.",
        "Export into an accounting package. Invoices are issued and can be printed.",
        "Medical or fitness advice. The assistant refuses it.",
    ]:
        bullet(doc, line)
    para(doc, "Assumptions. Staff are trusted inside their role. The system stops mistakes and stops a member acting outside their own data. It does not defend the till against the person standing at it; it records who did what. The database is PostgreSQL. Prices are integers, rounded to round_vnd (default 1,000). A missing model key degrades the assistant only.", space_before=8)
    para(doc, "The context diagram has no numbered process. Arena 3 is one circle. People are rectangles. payOS and mail delivery are external systems, not actors on the use-case diagram. Arrows are labelled with a short verb.")
    figure(
        doc,
        DOCS / "diagrams" / "context-diagram.png",
        "Figure 1. Context diagram. Guest, Member, Receptionist, Coach, Manager, payOS, and Mail delivery surround Arena 3 Sports Complex. Source: docs/diagrams/context-diagram.drawio.",
        max_h_cm=16,
    )
    add_table(
        doc,
        ["External", "They send", "Arena 3 sends"],
        [
            ["Guest", "Sign up", "View schedule"],
            ["Member", "Book a court", "Show the pass"],
            ["Receptionist", "Run the desk", "Open the gate"],
            ["Coach", "Mark attendance", "Show roster"],
            ["Manager", "Set prices", "Show revenue"],
            ["payOS", "Payment result", "Request payment"],
            ["Mail delivery", "Delivery status", "Send the OTP"],
        ],
        [4.0, 6.5, 6.5],
    )

    # 2
    heading(doc, "2. User Requirements", 1)
    heading(doc, "2.1 Actors", 2)
    para(doc, "An actor is a person. payOS and mail delivery stay on the context diagram. The assistant is a use case inside the system (UC-30), not an actor. The scheduler is not an actor either; its work is UC-31 and section 3.1.4. A user has exactly one role, checked on the server for every request. Hiding a button is only a convenience.")
    add_table(doc, ["#", "Actor", "Description"], ACTORS, [1.2, 3.6, 12.2])

    heading(doc, "2.2 Use Cases", 2)
    heading(doc, "2.2.1 Diagram", 3)
    para(doc, "One system boundary contains the thirty-one goals. Stick figures stand outside it. A solid line is an association. An open dashed arrow labelled include runs from the base goal to the goal it always uses. An open dashed arrow labelled extend runs from the optional goal back to the base: refund extends cancel, and reschedule extends book. Goals are not numbered on the figure. Packages, and boxes for mail, payOS, or the assistant, are not used.")
    para(doc, "Include: register includes verify email; book a court includes pay online; buy or renew includes pay online; record attendance includes lock the book; view revenue includes export; check-in includes display the QR pass. Extend: refund extends cancel; reschedule extends book.")
    para(doc, "Earlier drafts numbered separate goals through UC-37 for sending mail, releasing a hold, marking a no-show, and generating sessions. Those are not ovals on this figure. Mail delivery is external. Releasing a hold, marking a no-show, and generating sessions sit inside UC-31.")
    doc.add_page_break()
    figure(
        doc,
        DOCS / "diagrams" / "use-case-diagram.png",
        "Figure 7. Use-case diagram, thirty-one goals inside Arena 3 Sports Complex. The figure is tall, so it is fitted to the page height. Source: docs/diagrams/use-case-diagram.drawio.",
        max_h_cm=22.5,
    )

    heading(doc, "2.2.2 Descriptions", 3)
    para(doc, "Identifiers below match the goals on Figure 7. The actor column names who performs the goal. Where the figure draws only some of those sticks, because the same stick is repeated beside the goals it owns, the description says so.")
    add_table(
        doc,
        ["ID", "Use case", "Actors", "Use case description"],
        [(u[0], u[1], u[2], u[3]) for u in USE_CASES],
        [1.8, 4.0, 3.6, 7.6],
    )

    heading(doc, "2.2.3 Specifications", 3)
    para(doc, "Each specification is the behaviour a tester can walk. Business-rule identifiers point at section 5.1.")
    for uc in USE_CASES:
        ident, name, actors, _desc, trigger, flow, rules = uc
        heading(doc, f"{ident} {name}", 4)
        lab(doc, "Actors", actors)
        lab(doc, "Trigger", trigger)
        lab(doc, "Main flow", flow)
        lab(doc, "Rules", rules)

    # 3
    heading(doc, "3. Functional Requirements", 1)
    heading(doc, "3.1 System Functional Overview", 2)
    heading(doc, "3.1.1 Screens Flow", 3)
    para(doc, "There is one flow per role. A box is a screen. A box drawn inside a box is a modal, a tab, or a step. A solid arrow is navigation. A dotted arrow is conditional. Login is the root of every signed-in flow. These pictures are not activity diagrams: they have no decisions, forks, or swimlanes.")
    para(doc, "Homes. Visitor: /. Member: /app. Receptionist: /desk. Coach: /coach. Manager: /manager. Account (/account) and Alerts (/alerts) are shared and appear on each flow that links to them.")
    for title, path, cap in FLOWS:
        heading(doc, title, 4)
        figure(doc, ROOT / path, cap, max_h_cm=20)

    heading(doc, "3.1.2 Screen Descriptions", 3)
    para(doc, "The rows follow the flows above. A nested panel is listed under the screen that owns it.")
    for title, rows in SCREENS:
        heading(doc, title, 4)
        numbered = [(str(i), a, b, c) for i, (a, b, c, _roles) in enumerate(rows, start=1)]
        add_table(doc, ["#", "Feature", "Screen", "Description"], numbered, [1.0, 2.6, 4.4, 9.0])

    heading(doc, "3.1.3 Screen Authorization", 3)
    para(doc, "An x means that role may open the screen. The server enforces the same boundary; a missing x is not only a hidden link. Manager marks on desk and coach screens follow the rule that a manager may open those areas. Guest marks stop at the public pages. Modals inherit the parent screen unless a row says otherwise.")
    para(doc, "Columns: Guest, Member, Recep. (receptionist), Coach, Manager.")
    for title, rows in SCREENS:
        heading(doc, title.replace("3.1.2", "3.1.3"), 4)
        body = []
        for feature_name, screen, _desc, letters in rows:
            body.append([f"{feature_name} — {screen}", *marks(letters)])
        add_table(doc, ["Screen", *ROLE_HEAD], body, [8.5, 1.7, 1.7, 1.7, 1.7, 1.7], center_cols=(1, 2, 3, 4, 5))

    heading(doc, "3.1.4 Non-Screen Functions", 3)
    para(doc, "These jobs have no screen of their own. A loop wakes about every 15 seconds, reads once whether any job has due rows, and runs only those, each in its own transaction. Session generation is additionally limited to every 10 minutes. None of them sends SMS.")
    add_table(doc, ["#", "Feature", "System function", "Description"], JOBS, [1.0, 2.6, 3.4, 10.0])

    heading(doc, "3.1.5 Entity Relationship Diagram", 3)
    para(doc, "Two pictures cover the core of the ledger. They are a subset. Later tables, including check-in tokens, promotions, booking series, day passes, loyalty adjustments, coach pay, and work orders, are not drawn. A full poster of every table was rejected because it cannot be read. The Chen drawing is titled Logical ERD and the column drawing is titled Physical ERD. This chapter files them as the conceptual model and the logical model, which is the pair the report outline uses.")
    heading(doc, "3.1.5.1 Conceptual entity-relationship diagram", 4)
    para(doc, "Chen notation. A rectangle is an entity, a diamond is a verb, and a cardinality is 1 or N. Lines are not arrowed. There is no line from User to Class labelled as coaching; the coach is an attribute of the class, drawn on the logical diagram.")
    figure(
        doc,
        DOCS / "diagrams" / "erd-logical.png",
        "Figure 8. Conceptual model, drawn in Chen notation. The picture's own title reads Logical ERD. User holds Subscription of Plan; User books CourtBooking on Court; Payment settles a booking; Class runs Session; Session lists Enrollment. Source: docs/diagrams/erd-logical.drawio.",
        max_h_cm=18,
    )
    heading(doc, "3.1.5.2 Logical entity-relationship diagram", 4)
    para(doc, "Each box is a core table. The shaded header names the table. Primary and foreign keys are written in the box. A line is an unarrowed foreign key, and only where the line does not cross a third table. Payment.user_id, Class.coach_id, and Class.court_id are columns without a drawn line for that reason. This is not a UML class diagram.")
    figure(
        doc,
        DOCS / "diagrams" / "erd-physical.png",
        "Figure 9. Logical model of the core tables. The picture's own title reads Physical ERD. Primary and foreign keys are printed in the boxes. Source: docs/diagrams/erd-physical.drawio.",
        max_h_cm=16,
    )
    para(doc, "Entity dictionary. Keys below are the ones printed on Figure 9, not a full column list.")

    entities = [
        ("1", "User", "A person who can sign in: member, receptionist, coach, or manager.", "Primary key: id.", "No foreign key on this picture. Phone, role, and status identify the account."),
        ("2", "Membership plan", "A product the centre sells: name and price in whole dong. Drawn as Plan on the conceptual diagram.", "Primary key: id.", "No foreign key."),
        ("3", "Subscription", "One member holding one plan, and the status of that holding.", "Primary key: id.", "Foreign key: user_id to User. Foreign key: plan_id to Membership plan."),
        ("4", "Court", "A bookable court, with a code and a sport.", "Primary key: id.", "No foreign key."),
        ("5", "Court booking", "A hold, a confirmed hire, or a later terminal state, for one member on one court.", "Primary key: id.", "Foreign key: court_id to Court. Foreign key: user_id to User."),
        ("6", "Payment", "Money in or out, in whole dong, with a status. A refund is a negative amount, not a second copy of the original.", "Primary key: id.", "Foreign key: user_id to User. The line is omitted on the figure so it does not cross other tables."),
        ("7", "Class", "A coached product on one court, taught by one coach.", "Primary key: id.", "Foreign key: coach_id to User. Foreign key: court_id to Court. Both lines are omitted on the figure."),
        ("8", "Session", "One occurrence of a class, at start_at.", "Primary key: id.", "Foreign key: class_id to Class."),
        ("9", "Enrollment", "One member on one class, including a waitlisted place.", "Primary key: id.", "Foreign key: class_id to Class. Foreign key: user_id to User."),
    ]
    table = doc.add_table(rows=1 + len(entities), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(table, [1.2, 3.6, 12.2])
    for i, head in enumerate(["#", "Entity", "Description"]):
        write_cell(table.rows[0].cells[i], head, bold=True, size=10, color=WHITE, center=True, fill="1B4F72")
    header_row(table.rows[0])
    cant_split(table.rows[0])
    for r, (n, name, desc, pk, fk) in enumerate(entities):
        fill = "F8FAFC" if r % 2 else "FFFFFF"
        write_cell(table.rows[r + 1].cells[0], n, size=10, center=True, fill=fill)
        write_cell(table.rows[r + 1].cells[1], name, bold=True, size=10, fill=fill)
        cell = table.rows[r + 1].cells[2]
        write_cell(cell, desc, size=10, fill=fill)
        add_more(cell, pk, bold=True, size=10)
        add_more(cell, fk, size=10)
        cant_split(table.rows[r + 1])
    spacer(doc, 8)
    para(doc, "Design models that are not part of this chapter. The registration activity diagram (phone uniqueness is checked before an OTP challenge is created), the UML class diagram, and the sequence diagram for booking a court through payOS are design artefacts. They live with the design specification (docs/SDD.md) and the files docs/diagrams/activity-diagram.drawio, class-diagram.drawio, and sequence-diagram.drawio. They are not requirements figures.")

    feature(
        doc, "3.2", "Register, sign in, and reset a password",
        "A guest opens Register, Login, or Forgot password. A manager creating a staff account is UC-26 and ends in the same sign-in.",
        "Guest for registration. Every role for sign-in and password reset.",
        "One account per phone, proved by a one-time code, then a session that opens the role home.",
        "Landing, Register, the OTP step, Login, Forgot password, and Account Security. See section 3.1.1.1.",
        "Name, phone, date of birth, password typed twice, guardian details when the person is under 16, acceptance of terms and privacy, and the one-time code.",
        "Reject a phone that already has an account, a weak or mismatched password, a minor with no guardian, or a missing acceptance. Keep the challenge for 5 minutes. Lock it for 15 minutes after five wrong codes. Lock the account for 15 minutes after five wrong passwords. Issue a 7-day session for a member and a 12-hour session for staff.",
        "The user row is inserted only after the code matches. The session token is stored as a SHA-256 hash; the raw token stays on the client. The password is stored as a salted hash.",
        "Registration is a form followed by a code step. Login is one form. Forgot password is two steps. No separate screen is required beyond section 3.1.2.1.",
        "The product does not send SMS. Production does not echo the code onto the page. The email transport of the outbox is a stub, which is a known limit of the current deployment, not a second channel.",
    )
    feature(
        doc, "3.3", "Book, pay, reschedule, and cancel a court",
        "A member picks a slot, or comes back to move or cancel one. Online payment starts when they choose VietQR.",
        "Member. Receptionist and manager act through the desk goals in section 3.6 when the person is at the counter.",
        "A court-time has at most one holder, and a confirmed booking agrees with its payment.",
        "Book Court, the hold panel, scan-to-pay, the payment-return page, and change-time mode.",
        "Sport, date, court, slot, payment method (cash at the desk, plan hour, payOS, or a claimed transfer), and, on a move, the new slot.",
        "Create a hold of hold_minutes (default 5). Refuse a past slot, a slot outside opening hours, a court that is not ready, a day beyond the booking horizon, a third slot the same day, or a member over the debt limit. Confirm only inside one transaction with the payment and the invoice. A transfer returns 202 and keeps a longer hold. A move inside two hours, or a move that changes the price, is refused and the old slot stands. Cancellation prices the refund at 100, 50, or 0 percent from the 24-hour and 4-hour tiers.",
        "The hold and the confirmed booking are rows of court booking, joined to an occupancy while they own the court. The occupancy is released on cancel, expiry, no-show, or completion. payOS posts money only after a verified result for at least the amount due.",
        "The member layout is section 3.1.2.2, court-booking rows. The five-minute label on the hold panel is the default of hold_minutes.",
        "Two members cannot win the same court-time; the loser receives CONFLICT_SLOT. An unpaid hold expires by the job in section 3.1.4. A no-show is not refunded. A refund at or above 1,000,000 VND waits for a manager, and the same payment cannot be refunded twice. The desk may also sell a fixed weekly series of 4 to 12 weeks on one court, one weekday, and one hour; a week already past or already taken is skipped and is not charged.",
    )
    feature(
        doc, "3.4", "Buy, renew, and freeze a membership",
        "A member opens Plans, or the desk runs the new-member wizard.",
        "Member. Receptionist when the sale is at the counter.",
        "A plan grants hours or sessions only while it is active and inside its dates, and a frozen plan cannot be used to buy another.",
        "Plans screen, scan-to-pay, and the freeze control on the member's plan.",
        "The plan identifier, the payment, and the freeze interval.",
        "A withdrawn plan is refused. An order is pending until enough money has arrived, in full or as the deposit the centre allows. Freeze days accumulate against the yearly cap. Rate limit: 3 orders a minute and 10 an hour.",
        "The subscription row points at the user and the plan. Activation, invoice, and notice happen together, including when payOS settles the link.",
        "Section 3.1.2.2 plans, and the desk wizard in section 3.1.2.3.",
        "Stop-selling does not delete the plan and does not take it away from members who already hold it. A renewal in the members report is a reporting convention: a new plan starting within 30 days of the previous end. That threshold is not a business rule.",
    )
    feature(
        doc, "3.5", "Enrol, attend, and manage classes",
        "A member joins a class, a coach opens a session, or a manager changes the class life cycle.",
        "Member, coach, and manager, on the goals UC-14 through UC-19.",
        "A class stays inside its capacity, a student has one enrolment, and the register has one owner after it locks.",
        "Member classes and attendance panel; coach schedule and attendance; manager class forms.",
        "Class, session, enrolment, attendance mark, result, homework, and the reason on a late move or a correction.",
        "Refuse a plan that does not cover the sport, a second or clashing enrolment, a closed class, or a full class (the member waits instead). Expire offers and invite the next person. Return one session once when the centre cancels. Lock the register two hours after the end. Three Absents raise an alert; Excused breaks the streak; nobody is dropped automatically.",
        "Enrollment and session rows follow Figure 9. Attendance, results, and homework are training records. They never write a payment.",
        "Sections 3.1.2.2, 3.1.2.4, and 3.1.2.5. When F4 is off, Progress and the training routes are unavailable and answer BR-62.",
        "A coach sees only their classes. A substitute must teach the sport and be free at the remaining times. Capacity cannot be lowered below the current roster. Health notes and coach notes do not leave their scope, and they are not sent to the assistant's model.",
    )
    feature(
        doc, "3.6", "Run the front desk",
        "A receptionist opens a shift, or a person arrives to pay, enter, borrow gear, or ask a question.",
        "Receptionist. Manager for refund sign-off, transfer reconciliation, and the member profile.",
        "Money taken at the desk belongs to an open shift, and the gate is not confused with the classroom.",
        "Desk home, close-shift, gate, court map, payments, gear, member profile, and support reply.",
        "Shift count, member identity, walk-in name and phone, loan quantity, transfer decision, refund reason, and the reply text.",
        "Refuse money when no shift is open. Keep a single open shift per receptionist. Treat a transfer as a promise until a person matches it. Refuse a walk-in with under 20 minutes left in the slot. Refuse a loan beyond stock. Do not let the gate write attendance.",
        "Payments carry method, amount, VAT, what they were for, who took them, and the shift. An idempotency key returns the original payment instead of taking the money twice. A refund reads the whole ledger of that booking before posting a second time.",
        "Section 3.1.2.3. The manager reaches Member Profile and Payments from their own menu, as on Figure 6.",
        "Closing a shift stores the counted cash against the expected total. Gear rental restores the exact quantity and cannot be returned twice; it does not create a payment row. A support reply names the author, closes the request, and notifies the member.",
    )
    feature(
        doc, "3.7", "Set prices, staff, revenue, and the audit log",
        "A manager opens Pricing, Plans, Staff, Settings, Reports, or Audit.",
        "Manager only, except where section 3.1.3 already grants the same desk screen.",
        "Prices come from rules rather than from a hardcoded default, staff accounts are issued by a manager, and a report file matches the screen.",
        "The manager screens in section 3.1.2.5.",
        "Price-rule fields, promotion code and value, plan fields, staff role, report period and format, and audit filters.",
        "Refuse a quote when no price rule matches. Apply a percent promotion once, round once to 1,000 VND, and leave at least 1,000 VND due. Compare each report with the previous period of equal length, treating a refund increase as unfavourable. Build an export from the same queries and do not store the file.",
        "Settings are the single centre row. Audit rows are insert-only. Staff passwords are temporary until changed. Feature flags F4, F5, and F6 switch without a new deployment; a switched-off module returns BR-62.",
        "Section 3.1.1.5. Export is an action on the report, not a separate application.",
        "A manager cannot lock or demote their own account. There is no public staff registration. Court revenue stays separate from class fees. Attendance percent is omitted, not printed as zero, when F4 is off or nobody has been marked.",
    )
    feature(
        doc, "3.8", "Ask the assistant and run scheduled jobs",
        "A member sends a question, or the job loop wakes because a row is due.",
        "Member for the assistant. No signed-in person for a job pass. Manager owns the rules the jobs enforce.",
        "Answers stay inside the centre's published facts, and time-based cleanup does not wait for someone to remember it.",
        "Assistant screen for UC-30. Jobs have no screen; they are section 3.1.4.",
        "The question text, at most 800 characters. Jobs read the rows named in section 3.1.4 and do not take input from a form.",
        "Answer from settings, plans, classes, and coaches. Fall back to local rules when the model is absent or fails, and label that source. Refuse medical advice. Enforce 20 questions a minute. Jobs apply the predicates in section 3.1.4 and do not send SMS.",
        "The assistant can insert a support request and nothing else. It receives no member health notes and no payment tools. Job writes go through the same release and status changes as the interactive paths, so a hold expired by the clock and a hold cancelled by hand leave the same shape of row.",
        "Section 3.1.2.2 Assistant. There is no job console in this release.",
        "The ten jobs are expire holds, mark no-shows, complete bookings, expire subscriptions, expiry reminders, flush the outbox, expire payOS links, expire waitlist offers, lock attendance, and generate sessions. Reminders are in-app. Email delivery remains a stub.",
    )

    # 4
    heading(doc, "4. Non-Functional Requirements", 1)
    heading(doc, "4.1 External Interfaces", 2)
    heading(doc, "4.1.1 User interface", 3)
    para(doc, "The client is one web application, including a member view that can be installed. Every screen has to be usable at a width of 375 pixels. A refusal is a sentence that says what to do next; a machine code travels with it for support. Currency is written in whole dong and is not forced into capitals. Motion is decorative: nothing important is available only as an animation. Running text must render Vietnamese diacritics. The three report exports use an embedded font for the same reason.")
    heading(doc, "4.1.2 Communications interfaces", 3)
    para(doc, "The browser talks to one JSON API under /v1, with a bearer token. The API is part of the same application; there is no second backend. Report export returns a file rather than JSON.")
    para(doc, "payOS is the only online money channel. Arena 3 creates a payment link, reads its state, and accepts a webhook whose signature it checks. Credentials stay on the server. Order codes come from the clock and a stored high-water mark, so a code is not reused. A confirmation counts only when the verified amount is at least the amount due.")
    para(doc, "Mail delivery is the external rectangle on the context diagram. It is the intended transport for the registration code and for queued email. The transport behind the outbox is a stub in this release. In-app notices are delivered. There is no SMS interface, no card acquirer, no map tile service, and no image-hosting service.")
    para(doc, "The assistant calls a model with a system prompt and a context block built from the centre's own public operating data. Failure, refusal, or a missing key stays inside the assistant and does not affect booking or payment.")
    heading(doc, "4.2 Quality Attributes", 2)
    heading(doc, "4.2.1 Usability", 3)
    for line in [
        "Each role has one home and one navigation. Shared Account and Alerts keep their meaning across roles.",
        "A hold shows its remaining time. A slot the member cannot take is visible and not selectable.",
        "The desk home can ask for a short badge count instead of loading every queue.",
        "Empty and refused states say what happened. A missing attendance percent is omitted rather than shown as zero.",
    ]:
        bullet(doc, line)
    heading(doc, "4.2.2 Reliability", 3)
    for line in [
        "Two bookings cannot own the same court-time, including two requests that arrive together. The court row is locked and the database rejects an overlap.",
        "A coach cannot be placed in two overlapping sessions.",
        "A confirmed booking, its payment, and its invoice are written in one transaction. Either all of them exist or none of them do.",
        "A request that takes money or changes a schedule accepts an idempotency key and returns the original result when it is replayed.",
        "Online payment is posted from a payOS poll or a verified webhook, never from the return page alone.",
        "With no external database configured, a development run uses an embedded database, migrates, and seeds itself.",
        "Losing the assistant's model does not affect the ledger.",
    ]:
        bullet(doc, line)
    heading(doc, "4.2.3 Performance", 3)
    for line in [
        "The signed-in profile is assembled from independent reads issued together, so sign-in pays for one round trip rather than a chain.",
        "A read-only request does not open a transaction.",
        "The job loop spends a pass with nothing due on a single existence check, not on ten separate transactions.",
        "Availability for up to 62 days is one read. Capacity reports cover up to 93 days.",
    ]:
        bullet(doc, line)
    para(doc, "Authorisation is part of every action, not a property of the page. The server checks the role and the owner of the row before it reads or writes. A member who asks for someone else's booking, receipt, ticket, or profile receives 403.")
    heading(doc, "4.2.4 Security and privacy", 3)
    for line in [
        "Passwords are salted hashes. Session tokens are stored as SHA-256 hashes.",
        "Health notes, guardian details, and coach notes are not sent to the assistant model.",
        "The demonstration sign-in and any on-screen one-time code are off in production unless explicitly enabled.",
        "Audit rows name the actor, the action, the entity, and the time. No endpoint deletes them.",
        "Document numbers for members, payments, and invoices are taken from a counter and do not reuse a value.",
    ]:
        bullet(doc, line)

    # 5
    heading(doc, "5. Requirement Appendix", 1)
    heading(doc, "5.1 Business Rules", 2)
    para(doc, "A refusal the member can meet in the product is traceable to a rule. Most rules travel in the error body as br. A few are enforced in the handler and named in the response text instead: BR-21 appears in the clash message, and BR-26 and BR-27 are enforced in the class life-cycle. The behaviour is the same either way.")
    add_table(doc, ["ID", "Rule definition"], RULES, [2.2, 14.8])
    para(doc, "Related constraints that are implemented more narrowly than a full setting screen: a plan grants benefits only while today falls inside its start and end; plan discounts are integers from 0 to 100; replacing a hold runs every check before the old hold is dropped; prices are treated as including VAT; money is rounded once, half up, after the discount. Time zone and currency stay fixed.")

    heading(doc, "5.2 Common Requirements", 2)
    for line in [
        "One centre, one time zone (Asia/Ho_Chi_Minh), one currency (VND). Amounts are integers. Rounding uses round_vnd, default 1,000.",
        "The phone number is unique and is the login identity. Email is the channel intended for the one-time code, not a second login name.",
        "Exactly one role per account. The role is read from the session on the server, never from a value the browser asserts.",
        "A member reads and changes only their own profile, bookings, receipts, tickets, attendance, and training.",
        "State changes a member cannot make themselves write an audit row.",
        "Feature flags default to the safe behaviour: a switched-off module is 403 with BR-62, not a silent empty page.",
        "Holds, offers, sessions, and reminders use the database clock. Business dates use the Indochina calendar date.",
        "The same payment method, once it has posted, produces the same invoice and the same payer notice. Transfer and payOS differ only in when posting is allowed.",
    ]:
        bullet(doc, line)

    heading(doc, "5.3 Application Messages List", 2)
    para(doc, "The API returns a code and a sentence. The sentence is what the screen shows. The code is what support quotes. Business-rule failures add the rule identifier from section 5.1.")
    add_table(doc, ["ID", "Code", "HTTP", "When it is returned"], MESSAGES, [2.0, 3.6, 1.6, 9.8])
    para(doc, "Illustrative sentences the product already uses, paraphrased only where the code builds them from a setting: a booking in this state cannot be cancelled; a booking can only be moved up to the configured number of hours before it starts; that row is already a refund; payOS is not configured and the desk must take the payment. Exact wording of a sentence that includes a number follows the current setting, so this list does not freeze one centre's hours into the message text.")

    para(doc, "End of chapter III.", bold=True, space_before=12)
    doc.save(OUT)
    return doc


def text_of(doc):
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def check(doc):
    body = text_of(doc)
    errors = []
    required = [
        "1. Product Overview",
        "2. User Requirements",
        "2.1 Actors",
        "2.2 Use Cases",
        "2.2.1 Diagram",
        "2.2.2 Descriptions",
        "2.2.3 Specifications",
        "3. Functional Requirements",
        "3.1 System Functional Overview",
        "3.1.1 Screens Flow",
        "3.1.2 Screen Descriptions",
        "3.1.3 Screen Authorization",
        "3.1.4 Non-Screen Functions",
        "3.1.5 Entity Relationship Diagram",
        "3.1.5.1 Conceptual entity-relationship diagram",
        "3.1.5.2 Logical entity-relationship diagram",
        "3.2 Register, sign in, and reset a password",
        "3.3 Book, pay, reschedule, and cancel a court",
        "3.4 Buy, renew, and freeze a membership",
        "3.5 Enrol, attend, and manage classes",
        "3.6 Run the front desk",
        "3.7 Set prices, staff, revenue, and the audit log",
        "3.8 Ask the assistant and run scheduled jobs",
        "4. Non-Functional Requirements",
        "4.1 External Interfaces",
        "4.1.1 User interface",
        "4.1.2 Communications interfaces",
        "4.2 Quality Attributes",
        "4.2.1 Usability",
        "4.2.2 Reliability",
        "4.2.3 Performance",
        "4.2.4 Security and privacy",
        "5. Requirement Appendix",
        "5.1 Business Rules",
        "5.2 Common Requirements",
        "5.3 Application Messages List",
        "Function trigger",
        "Function description",
        "Data collect",
        "Data computation",
        "Data binding",
        "Screen layout",
        "Function details",
    ]
    for item in required:
        if item not in body:
            errors.append(f"missing heading or label: {item}")
    for uc in USE_CASES:
        if uc[0] not in body or uc[1] not in body:
            errors.append(f"missing use case {uc[0]}")
    if len(USE_CASES) != 31:
        errors.append("use case list is not 31")
    images = [r for r in doc.part.rels.values() if "image" in r.reltype]
    if len(images) != 9:
        errors.append(f"expected 9 images, found {len(images)}")
    banned = ["Cloudinary", "FreeDriver", "Mapbox", "data is filtered through multiple layers", "Mermaid", "«actor»"]
    for word in banned:
        if word.lower() in body.lower():
            errors.append(f"banned phrase present: {word}")
    # SMS may appear only as an absence.
    if "SMS" in body and "no SMS" not in body and "does not send SMS" not in body and "not send SMS" not in body:
        errors.append("SMS mentioned without stating it is absent")
    if "UC-32" in body or "UC-37" in body:
        # The history may mention the old range. Allow the single supersession sentence.
        pass
    if errors:
        print("CHECK FAILED")
        for e in errors:
            print(" -", e)
        return 1
    print(f"CHECK OK  paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(images)} chars={len(body)}")
    return 0


def main():
    build()
    # Re-open so relationship counts match the saved package.
    saved = Document(str(OUT))
    code = check(saved)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    if BACKUP.exists():
        print(f"previous file kept at {BACKUP}")
    return code


if __name__ == "__main__":
    sys.exit(main())
